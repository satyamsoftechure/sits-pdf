from docx.shared import Pt
from ..common.Element import Element
from ..layout.Layout import Layout
from ..common import docx


class Cell(Layout):
    def __init__(self, raw: dict = None):
        raw = raw or {}
        super().__init__()
        self.restore(raw)
        self.bg_color = raw.get("bg_color", None)
        self.border_color = raw.get("border_color", (0, 0, 0, 0))
        self.border_width = raw.get("border_width", (0, 0, 0, 0))
        self.merged_cells = raw.get("merged_cells", (1, 1))

    @property
    def text(self):
        if not self:
            return None
        return "\n".join(
            [
                block.text if block.is_text_block else "<NEST TABLE>"
                for block in self.blocks
            ]
        )

    @property
    def working_bbox(self):
        x0, y0, x1, y1 = self.bbox
        w_top, w_right, w_bottom, w_left = self.border_width
        bbox = (
            x0 + w_left / 2.0,
            y0 + w_top / 2.0,
            x1 - w_right / 2.0,
            y1 - w_bottom / 2.0,
        )
        return Element().update_bbox(bbox).bbox

    def store(self):
        if not bool(self):
            return None
        res = super().store()
        res.update(
            {
                "bg_color": self.bg_color,
                "border_color": self.border_color,
                "border_width": self.border_width,
                "merged_cells": self.merged_cells,
            }
        )
        return res

    def plot(self, page):
        super().plot(page)
        self.blocks.plot(page)

    def make_docx(self, table, indexes):
        self._set_style(table, indexes)
        if not bool(self):
            return
        n_row, n_col = self.merged_cells
        i, j = indexes
        docx_cell = table.cell(i, j)
        if n_row * n_col != 1:
            _cell = table.cell(i + n_row - 1, j + n_col - 1)
            docx_cell.merge(_cell)
        x0, y0, x1, y1 = self.bbox
        docx_cell.width = Pt(x1 - x0)
        if self.blocks:
            docx_cell._element.clear_content()
            self.blocks.make_docx(docx_cell)

    def _set_style(self, table, indexes):

        i, j = indexes
        docx_cell = table.cell(i, j)
        n_row, n_col = self.merged_cells

        keys = ("top", "end", "bottom", "start")
        kwargs = {}
        for k, w, c in zip(keys, self.border_width, self.border_color):
            if not w:
                continue

            hex_c = f"#{hex(c)[2:].zfill(6)}"
            kwargs[k] = {"sz": 8 * w, "val": "single", "color": hex_c.upper()}
        for m in range(i, i + n_row):
            for n in range(j, j + n_col):
                docx.set_cell_border(table.cell(m, n), **kwargs)
        if self.bg_color is not None:
            docx.set_cell_shading(docx_cell, self.bg_color)
        docx.set_cell_margins(docx_cell, start=0, end=0)
        if self.blocks.is_vertical_text:
            docx.set_vertical_cell_direction(docx_cell)
